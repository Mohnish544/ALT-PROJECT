"""
Text Extractor
==============
Extracts raw text from PDF, DOCX, and TXT files.
Handles multi-column PDFs via pdfplumber, with pypdf as fallback.
"""
from __future__ import annotations

import logging
from pathlib import Path
from typing import NamedTuple

import chardet
import pdfplumber
import pypdf
from docx import Document as DocxDocument

logger = logging.getLogger(__name__)


class ExtractionResult(NamedTuple):
    text: str
    page_count: int
    file_type: str
    char_count: int
    extraction_method: str


def extract(file_path: Path) -> ExtractionResult:
    """
    Extract plain text from a supported file type.

    Args:
        file_path: Absolute path to the uploaded file.

    Returns:
        ExtractionResult with the raw text and metadata.

    Raises:
        ValueError: If the file type is unsupported or extraction fails.
    """
    suffix = file_path.suffix.lower()

    if suffix == ".pdf":
        return _extract_pdf(file_path)
    elif suffix == ".docx":
        return _extract_docx(file_path)
    elif suffix == ".txt":
        return _extract_txt(file_path)
    else:
        raise ValueError(
            f"Unsupported file type: {suffix!r}. Accepted: .pdf, .docx, .txt"
        )


# ── PDF Extraction ────────────────────────────────────────────────────────────

def _extract_pdf(path: Path) -> ExtractionResult:
    """
    Try pdfplumber first (better for multi-column, tables).
    Fall back to pypdf if pdfplumber yields empty text.
    """
    text, page_count, method = _pdfplumber_extract(path)

    if not text.strip():
        logger.warning("pdfplumber returned empty text; falling back to pypdf")
        text, page_count = _pypdf_extract(path)
        method = "pypdf"

    if not text.strip():
        raise ValueError(
            f"Could not extract any text from {path.name}. "
            "The PDF may be scanned/image-based. Please use a text-based PDF."
        )

    return ExtractionResult(
        text=text,
        page_count=page_count,
        file_type="pdf",
        char_count=len(text),
        extraction_method=method,
    )


def _pdfplumber_extract(path: Path) -> tuple[str, int, str]:
    pages: list[str] = []
    with pdfplumber.open(path) as pdf:
        page_count = len(pdf.pages)
        for page in pdf.pages:
            # Extract text; extract_tables and combine if present
            page_text = page.extract_text(x_tolerance=3, y_tolerance=3) or ""

            # Also extract tables and append as pipe-delimited text
            for table in page.extract_tables():
                for row in table:
                    cleaned = [str(cell).strip() if cell else "" for cell in row]
                    page_text += "\n" + " | ".join(cleaned)

            pages.append(page_text)

    return "\n\n".join(pages), page_count, "pdfplumber"


def _pypdf_extract(path: Path) -> tuple[str, int]:
    reader = pypdf.PdfReader(str(path))
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n\n".join(pages), len(reader.pages)


# ── DOCX Extraction ───────────────────────────────────────────────────────────

def _extract_docx(path: Path) -> ExtractionResult:
    doc = DocxDocument(str(path))
    paragraphs = [para.text for para in doc.paragraphs if para.text.strip()]

    # Also extract table cell text
    for table in doc.tables:
        for row in table.rows:
            row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_texts:
                paragraphs.append(" | ".join(row_texts))

    text = "\n".join(paragraphs)
    if not text.strip():
        raise ValueError(f"No text found in {path.name}. The DOCX may be empty.")

    # DOCX doesn't have a reliable page count; estimate at 300 words/page
    word_count = len(text.split())
    estimated_pages = max(1, word_count // 300)

    return ExtractionResult(
        text=text,
        page_count=estimated_pages,
        file_type="docx",
        char_count=len(text),
        extraction_method="python-docx",
    )


# ── TXT Extraction ────────────────────────────────────────────────────────────

def _extract_txt(path: Path) -> ExtractionResult:
    raw_bytes = path.read_bytes()

    # Auto-detect encoding
    detected = chardet.detect(raw_bytes)
    encoding = detected.get("encoding") or "utf-8"

    try:
        text = raw_bytes.decode(encoding)
    except (UnicodeDecodeError, LookupError):
        text = raw_bytes.decode("utf-8", errors="replace")
        logger.warning("Fell back to UTF-8 with replacement for %s", path.name)

    if not text.strip():
        raise ValueError(f"The text file {path.name} is empty.")

    word_count = len(text.split())
    estimated_pages = max(1, word_count // 300)

    return ExtractionResult(
        text=text,
        page_count=estimated_pages,
        file_type="txt",
        char_count=len(text),
        extraction_method="chardet+utf-8",
    )
